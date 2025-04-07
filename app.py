# --- START OF FILE app.py ---

import re
import os
import requests
import logging
from datetime import datetime
from flask import Flask, render_template, request, url_for # Keep url_for if used elsewhere, maybe not needed now
from werkzeug.utils import secure_filename
from dotenv import load_dotenv

# --- spaCy Import and Model Loading ---
try:
    import spacy
    try:
        nlp = spacy.load("en_core_web_sm")
        print("spaCy model 'en_core_web_sm' loaded successfully.")
    except OSError:
        print("spaCy model 'en_core_web_sm' not found.")
        print("Please run: python -m spacy download en_core_web_sm")
        nlp = None
except ImportError:
    print("WARNING: spacy library not installed. NER detection will be disabled.")
    print("Install using: pip install spacy")
    nlp = None
# --- End spaCy Setup ---

load_dotenv() # Load variables from .env file

# Import file reading libraries
try:
    import PyPDF2
except ImportError:
    print("WARNING: PyPDF2 library not installed. PDF processing will be disabled.")
    print("Install using: pip install PyPDF2")
    PyPDF2 = None

try:
    import docx
except ImportError:
    print("WARNING: python-docx library not installed. DOCX processing will be disabled.")
    print("Install using: pip install python-docx")
    docx = None

app = Flask(__name__)

# --- Audit Logging Setup ---
audit_logger = logging.getLogger('audit')
audit_logger.setLevel(logging.INFO)
audit_handler = logging.FileHandler('audit.log', encoding='utf-8')
audit_handler.setLevel(logging.INFO)
audit_formatter = logging.Formatter('%(asctime)s | %(message)s', datefmt='%Y-%m-%d %H:%M:%S')
audit_handler.setFormatter(audit_formatter)
if not audit_logger.hasHandlers():
    audit_logger.addHandler(audit_handler)
# --- End Audit Logging Setup ---

# --- Configuration ---
ALLOWED_EXTENSIONS = {'txt', 'pdf', 'docx'}
LOG_FILE_PATH = 'audit.log' # Define log path centrally

def allowed_file(filename):
    """Checks if the filename has an allowed extension."""
    return '.' in filename and filename.rsplit('.', 1)[1].lower() in ALLOWED_EXTENSIONS

# --- Text Extraction Helper Functions ---
def extract_text_from_txt(file_stream):
    try: file_stream.seek(0); return file_stream.read().decode('utf-8', errors='ignore')
    except Exception as e: print(f"Error reading TXT file: {e}"); return None

# *** CORRECTED extract_text_from_pdf FUNCTION ***
def extract_text_from_pdf(file_stream):
    """Extracts text from a PDF file stream using PyPDF2."""
    if not PyPDF2:
        print("PyPDF2 library not installed. Cannot process PDF.")
        return None
    text = ""
    try:
        file_stream.seek(0)
        reader = PyPDF2.PdfReader(file_stream)
        if reader.is_encrypted:
            print("PDF is encrypted, cannot extract text.") # More informative print
            return "[ERROR: PDF is encrypted]" # Return the error marker

        # This loop should only run if the PDF is *not* encrypted
        for page in reader.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n" # Correct indentation

        return text if text else None # Return None if no text extracted

    except PyPDF2.errors.PdfReadError as e:
        print(f"Error reading PDF file (PdfReadError): {e}.")
        return None
    except Exception as e:
        print(f"Error reading PDF file: {e}")
        return None
# *** END CORRECTION ***

def extract_text_from_docx(file_stream):
    if not docx: return None
    text = "";
    try:
        file_stream.seek(0); document = docx.Document(file_stream)
        for para in document.paragraphs: text += para.text + "\n"
        return text if text else None
    except Exception as e: print(f"Error reading DOCX file: {e}"); return None

# --- Hugging Face API Setup ---
HF_API_KEY = os.environ.get("HF_API_KEY")

if not HF_API_KEY:
    print("WARNING: HF_API_KEY environment variable not set. LLM functionality will be disabled.")

API_URL = "https://api-inference.huggingface.co/models/distilbert-base-uncased-finetuned-sst-2-english"
headers = {"Authorization": f"Bearer {HF_API_KEY}"}

# --- LLM Query Function ---
def query_llm(text):
    if not HF_API_KEY or not API_URL: return "[LLM Query Skipped - API Key or URL not configured]"
    if not text or not isinstance(text, str) or not text.strip(): return "[LLM Query Skipped - Input text is empty or invalid]"
    if "distilbert" in API_URL or "sentiment" in API_URL: prompt = text; payload = {"inputs": prompt}; is_classification = True
    else: prompt = f"Process the following text and provide a relevant response:\n\n---\n{text}\n---\n\nResponse:"; payload = {"inputs": prompt, "parameters": {"max_new_tokens": 250, "return_full_text": False, "temperature": 0.7, "top_p": 0.9,}}; is_classification = False
    try:
        print(f"Sending request to LLM: {API_URL}"); response = requests.post(API_URL, headers=headers, json=payload, timeout=90)
        response.raise_for_status(); result = response.json(); print(f"LLM Raw Response: {result}"); generated_text = None
        if is_classification:
            if isinstance(result, list) and result and isinstance(result[0], list) and result[0]:
                if isinstance(result[0][0], dict) and 'label' in result[0][0]: top_result = result[0][0]; generated_text = f"Sentiment Analysis: Label='{top_result.get('label')}', Score={top_result.get('score'):.4f}"
                else: generated_text = f"[LLM Response (Unknown Format): {str(result)[:200]}...]"
            else: generated_text = f"[LLM Response (Non-List Format): {str(result)[:200]}...]"
        else:
            if isinstance(result, list) and result:
                if isinstance(result[0], dict) and 'generated_text' in result[0]: generated_text = result[0]['generated_text']
                elif isinstance(result[0], str): generated_text = result[0]
            elif isinstance(result, dict) and 'generated_text' in result: generated_text = result.get('generated_text')
            if generated_text:
                if generated_text.strip().startswith(prompt.strip()) and not is_classification: generated_text = generated_text[len(prompt):].strip()
                elif "Response:" in generated_text: generated_text = generated_text.split("Response:", 1)[-1].strip()
        if generated_text: return generated_text.strip()
        else:
            print(f"Could not parse expected output from LLM response: {result}")
            if isinstance(result, dict) and 'error' in result:
                 api_error = result['error']; print(f"Hugging Face API Error Message: {api_error}")
                 if isinstance(api_error, str) and "currently loading" in api_error.lower(): return f"[LLM Info: Model is currently loading, please wait ({result.get('estimated_time', 'unknown'):.1f}s estimated) and try again.]"
                 else: return f"[Error: Received error from API: {api_error}]"
            return "[Error: LLM response format unexpected or empty]"
    except requests.exceptions.Timeout: print(f"Error querying LLM: Request timed out."); return "[Error: LLM request timed out]"
    except requests.exceptions.RequestException as e:
        error_detail = str(e); status_code = "N/A"; resp_text = "N/A";
        if e.response is not None: status_code = e.response.status_code; resp_text = e.response.text[:500]
        print(f"Error querying LLM ({API_URL}): {error_detail} | Status Code: {status_code} | Response: {resp_text}")
        if status_code == 401: return "[Error querying LLM: Authorization failed (401). Check your HF_API_KEY.]"
        elif status_code == 402: return f"[Error querying LLM: Payment Required (402). Access to model {API_URL.split('/')[-1]} requires a paid plan.]"
        elif status_code == 403: return f"[Error querying LLM: Forbidden (403). Ensure you accepted the terms for model {API_URL.split('/')[-1]} on Hugging Face.]"
        elif status_code == 429: return "[Error querying LLM: Rate limit possibly exceeded (429).]"
        elif status_code == 503: return "[Error querying LLM: Service Unavailable (503). Model might be loading or temporarily down. Try again later.]"
        elif status_code >= 500: return f"[Error querying LLM: Server error ({status_code}). Try again later.]"
        else: return f"[Error querying LLM: {str(e)}]"
    except (IndexError, KeyError, TypeError) as e: print(f"Error parsing LLM response: {e}. Response was: {result if 'result' in locals() else 'Not available'}"); return "[Error: Invalid response structure from LLM]"
    except Exception as e: print(f"An unexpected error occurred during LLM query: {e}"); return f"[An unexpected error occurred during LLM query: {str(e)}]"


# --- Redaction Functions ---
PLACEHOLDER_MAP = { # Full map here
    "Email Address": "[EMAIL]", "Phone Number Format (US)": "[PHONE]", "SSN Format (US)": "[SSN]",
    "Potential SSN (Loose Format)": "[SSN_LIKE]", "Potential Credit Card": "[CREDIT_CARD]",
    "API Key Format": "[API_KEY]", "Private Key Block": "[PRIVATE_KEY]", "Internal/Local IP": "[IP_ADDRESS]",
    "UUID Format": "[UUID]", "Secret Assignment": "[SECRET_VALUE]", "Password Assignment": "[PASSWORD_VALUE]",
    "Credential Assignment": "[CREDENTIAL_VALUE]", "Username Assignment": "[USERNAME_VALUE]",
    "Bank Account Number": "[BANK_ACCOUNT]", "Date Format (MM/DD/YY or YYYY)": "[DATE]",
    "Project Codename": "[PROJECT_CODENAME]", "Potential Generic Key/Token": "[GENERIC_KEY]",
    "JSON Key Pair": "[JSON_SECRET_PAIR]", "Value in Parentheses": "[SECRET_IN_PARENS]",
    "Value near Keyword": "[POTENTIAL_VALUE]", "PERSON": "[PERSON_NAME]", "GPE": "[LOCATION]",
    "LOC": "[LOCATION]", "ORG": "[ORGANIZATION]", "DATE": "[DATE_ENTITY]", "CARDINAL": "[NUMBER]",
    "MONEY": "[MONETARY_VALUE]", "NORP": "[GROUP]", "FAC": "[FACILITY]", "SSN Keyword": "[SSN_MENTIONED]",
    "DOB Keyword": "[DOB_MENTIONED]", "Credentials": "[CREDENTIALS_INFO]", "Personal Data (PII)": "[PII_INFO]",
    "Potential PII/Contextual": "[CONTEXTUAL_INFO]", "Confidential/Secret": "[CONFIDENTIAL_INFO]",
    "Legal": "[LEGAL_INFO]", "Intellectual Property": "[IP_INFO]", "Network": "[NETWORK_INFO]",
    "Identifier": "[IDENTIFIER_INFO]", "Keyword": "[SENSITIVE_KEYWORD]",
}
def get_placeholder(finding): # Logic remains same
    ftype = finding.get('type', ''); fcategory = finding.get('category', '')
    if ftype in PLACEHOLDER_MAP: return PLACEHOLDER_MAP[ftype]
    if fcategory in PLACEHOLDER_MAP: return PLACEHOLDER_MAP[fcategory]
    if ftype == 'Keyword': return PLACEHOLDER_MAP.get('Keyword', '[KEYWORD]')
    sanitized_category = re.sub(r'[^\w\s-]', '', fcategory).strip().upper().replace(' ','_')
    return f"[{sanitized_category or 'INFO'}]"
def redact_text(text, findings): # Logic remains same
    if not findings or not text: return text
    sorted_findings = sorted(findings, key=lambda f: (f['span'][0], -f['span'][1]))
    redacted = ""; last_end = 0; processed_indices = set()
    for finding in sorted_findings:
        start, end = finding['span']
        if not (0 <= start < end <= len(text)): continue
        if any(i in processed_indices for i in range(start, end)): continue
        if start >= last_end: redacted += text[last_end:start]
        else: print(f"Warning: Adjusting start due to potential finding overlap issue. Finding: {finding}")
        placeholder = get_placeholder(finding); redacted += placeholder; last_end = end
        for i in range(start, end): processed_indices.add(i)
    if last_end < len(text): redacted += text[last_end:]
    return redacted

# --- Confidentiality Check Function ---
def perform_confidentiality_check(text):
    if not text or not isinstance(text, str): return {"found_issues": False, "findings": []}
    findings = []; processed_text_positions = set()
    original_text = text; lower_text = original_text.lower()
    def add_finding(category, type, value, description, start_index, end_index):
        if not (0 <= start_index < end_index <= len(original_text)): print(f"DEBUG SKIP - Invalid span ({start_index},{end_index}) for value '{value[:50]}...'"); return
        span = (start_index, end_index)
        if any(i in processed_text_positions for i in range(start_index, end_index)): print(f"DEBUG SKIP - Overlap detected for {type} at {span}"); return
        print(f"DEBUG ADD - Adding Finding: {type} | Cat: {category} | Val: '{value}' | Span: {span}"); findings.append({"category": category, "type": type, "value": value, "description": description, "span": span})
        for i in range(start_index, end_index): processed_text_positions.add(i)
    # 1. spaCy NER Processing
    spacy_findings_temp = []
    if nlp:
        try:
            doc = nlp(original_text)
            for ent in doc.ents:
                label = ent.label_; category = "Unknown"; finding_type = label; entity_text = ent.text; entity_text_lower = entity_text.lower()
                if label in ["ORG", "MISC", "PRODUCT"] :
                    if entity_text_lower == "ssn": label = "SSN_KW_OVERRIDE"; category = "Personal Data (PII)"; finding_type = "SSN Keyword"
                    elif entity_text_lower == "dob": label = "DOB_KW_OVERRIDE"; category = "Personal Data (PII)"; finding_type = "DOB Keyword"
                if label == "PERSON": category = "Personal Data (PII)"; finding_type = "PERSON"
                elif label in ["GPE", "LOC"]: category = "Contextual"; finding_type = "LOCATION"
                elif label == "ORG": category = "Contextual"; finding_type = "ORGANIZATION"
                elif label == "DATE": category = "Contextual"; finding_type = "DATE_ENTITY"
                elif label in ["CARDINAL", "MONEY", "QUANTITY"]: category = "Potential PII/Contextual"; finding_type = label
                elif label in ["NORP", "FAC", "PRODUCT", "EVENT", "WORK_OF_ART", "LAW", "LANGUAGE"]: category = "Contextual"; finding_type = label
                if category != "Unknown": spacy_findings_temp.append({"category": category, "type": finding_type, "value": entity_text, "description": f"Detected '{label}' entity (via spaCy NER).", "span": (ent.start_char, ent.end_char)})
        except Exception as e: print(f"Error during spaCy processing: {e}")
    else: print("spaCy NER skipped as model is not loaded.")
    print(f"DEBUG: spaCy Findings Temp = {spacy_findings_temp}")
    # 2. Keyword Scanning
    keyword_categories = {
        "Legal": [r"attorney-client\s+privilege", r"legal\s+hold", r"litigation", r"confidential\s+settlement", r"under\s+seal", r"privileged\s+and\s+confidential", r"cease\s+and\s+desist", r"nda", r"non-disclosure\s+agreement"],
        "Confidential/Secret": [r"confidential", r"proprietary", r"secrets?", r"internal\s+use\s+only", r"trade\s+secret", r"classified", r"sensitive", r"do\s+not\s+distribute", r"private", r"restricted", r"not\s+for\s+public\s+release", r"passwords?", r"secret keys?", r"api keys?", r"credentials?", r"tokens?"],
        "Intellectual Property": [r"patent\s+pending", r"patent\s+application", r"trademark", r"copyright", r"invention\s+disclosure", r"prototype", r"roadmap", r"research\s+findings", r"algorithm", r"proprietary\s+algorithm"],
        "Personal Data (PII)": [r"dob", r"date\s+of\s+birth", r"passport\s+number", r"driver'?s\s+license", r"address", r"ssn", r"social security numbers?", r"credit cards?", r"bank accounts?", r"email addresses?", r"phone numbers?"],
    }
    temp_keyword_findings = []; keyword_processed_positions_check = set()
    if lower_text:
        for category, keywords in keyword_categories.items():
            for keyword_pattern in keywords:
                try:
                    for match in re.finditer(r'\b' + keyword_pattern + r'\b', lower_text, re.IGNORECASE):
                        start, end = match.start(), match.end()
                        if not any(pos in keyword_processed_positions_check for pos in range(start, end)):
                             temp_keyword_findings.append({"category": category, "type": "Keyword", "value": original_text[start:end], "description": f"Detected keyword: '{original_text[start:end]}'.", "span": (start, end)})
                             for i in range(start, end): keyword_processed_positions_check.add(i)
                except re.error as e: print(f"Warning: Skipping invalid regex keyword: '{keyword_pattern}' - {e}")
    print(f"DEBUG: Keyword Findings Temp = {temp_keyword_findings}")
    # 3. Keyword Proximity Scan
    KEYWORDS_INDICATING_VALUE = {'password', 'pass', 'pwd', 'secret', 'key', 'token', 'credential', 'credentials', 'username', 'user', 'login', 'userid', 'user_id', 'account', 'acct', 'accounts', 'ssn', 'social security number', 'credit card', 'bank account', 'email', 'phone'}
    POTENTIAL_VALUE_REGEX = r'\b([a-zA-Z0-9\-_+=/.@]{5,})\b'; PROXIMITY_WINDOW = 50
    potential_value_findings = []; proximity_checked_indices = set()
    for kw_finding in temp_keyword_findings:
        keyword_value_lower = kw_finding['value'].lower().rstrip('s')
        if keyword_value_lower in KEYWORDS_INDICATING_VALUE:
             kw_start, kw_end = kw_finding['span']; search_start = kw_end; search_end = min(kw_end + PROXIMITY_WINDOW, len(original_text)); search_text = original_text[search_start:search_end]
             try:
                  for value_match in re.finditer(POTENTIAL_VALUE_REGEX, search_text):
                       val_start_in_slice, val_end_in_slice = value_match.span(); val_start_abs = search_start + val_start_in_slice; val_end_abs = search_start + val_end_in_slice; value = value_match.group(1)
                       if value.lower() in ['is', 'are', 'was', 'the', 'a', 'an', 'and', 'for', 'my', 'your', 'number'] or (value.isdigit() and len(value)<3): continue
                       if not any(i in proximity_checked_indices for i in range(val_start_abs, val_end_abs)):
                            potential_value_findings.append({"category": kw_finding['category'], "type": "Value near Keyword", "value": value, "description": f"Potential value found near keyword '{kw_finding['value']}'.", "span": (val_start_abs, val_end_abs)})
                            for i in range(val_start_abs, val_end_abs): proximity_checked_indices.add(i)
             except re.error as e: print(f"Warning: Error during proximity value search: {e}")
    print(f"DEBUG: Proximity Value Findings Temp = {potential_value_findings}")
    # 4. Regex Pattern Matching
    regex_patterns = [
        (r'\b(password|passwd|secret|pwd|pass)\b(?:\s+(?:is|was|are|be)\s+|\s*[:=]\s*|\s+)(\S+)', "Password Assignment", "Credentials", "Potential password assignment.", re.IGNORECASE),
        (r'\b(api[\._]?key|access[\._]?key|secret[\._]?key|token|credential|auth_token)\b\s*[:=]\s*(?:\{|"|\'|)([^\s;\'"\}]+)(?:\}|"|\'|;)', "Credential Assignment", "Credentials", "Potential hardcoded key/token assignment in code.", re.IGNORECASE),
        (r'\b(user(?: |_)name|user|login|user_?id)\s*[:=]\s*\S+', "Username Assignment", "Credentials", "Potential username or login ID assignment.", re.IGNORECASE),
        (r'\b(secrets?|keys?|tokens?|passwords?|credentials?|values?)\b\s*\(([a-zA-Z0-9\-_+=]{6,})\)', "Value in Parentheses", "Credentials", "Potential secret value enclosed in parentheses after keyword.", re.IGNORECASE),
        (r'\b(sk_live|pk_live|rk_live|sk_test|pk_test|rk_test)_[0-9a-zA-Z]{24,}\b', "API Key Format", "Credentials", "Common API key format (e.g., Stripe)."),
        (r'-----BEGIN (RSA|OPENSSH|PGP|DSA|EC) PRIVATE KEY-----.*?-----END \1 PRIVATE KEY-----', "Private Key Block", "Credentials", "Private key block detected.", re.DOTALL | re.IGNORECASE),
        (r'\b[a-zA-Z0-9\+\/]{40,}\b', "Potential Generic Key/Token", "Credentials", "High entropy string (potential key/token)."),
        (r'"(?:access_key|secret_key|api_key|token)"\s*:\s*"([^"]+)"', "JSON Key Pair", "Credentials", "Potential sensitive keys in JSON.", re.IGNORECASE),
        (r'\b(ssn|social security(?: number)?)\s*(?:is|:|=)?\s*(\d{3}-?\d{2}-?\d{4})\b', "SSN Format (US)", "Personal Data (PII)", "US SSN pattern with indicator.", re.IGNORECASE),
        (r'(?<!\d-)\b(\d{3}-\d{2}-\d{4})\b(?!\d)', "SSN Format (US)", "Personal Data (PII)", "US SSN pattern (standalone)."),
        (r'(?<!\d-)\b(\d{7}|\d{9})\b(?!\d)', "Potential SSN (Loose Format)", "Personal Data (PII)", "7 or 9 digits (standalone, potential SSN)."),
        (r'\b(credit card|cc|card number|cc#|card no)\s*(?:is|:|=)?\s*(\d[\d -]{11,18}\d)\b', "Potential Credit Card", "Personal Data (PII)", "Potential Credit Card Number with indicator.", re.IGNORECASE),
        (r'\b(\d{4}[ -]?\d{4}[ -]?\d{4}[ -]?\d{4})\b', "Potential Credit Card", "Personal Data (PII)", "16 digits in groups of 4."),
        (r'\b(bank account(?: number| no)?|account no|account number|acct no)\s*(?:is|:|=|,?\s+which\s+is)?\s*(\d[\d\s-]{5,}\d?)', "Bank Account Number", "Personal Data (PII)", "Potential bank account number.", re.IGNORECASE),
        (r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b', "Email Address", "Personal Data (PII)", "Email address format."),
        (r'\(?\b\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}\b', "Phone Number Format (US)", "Personal Data (PII)", "US phone number format."),
        (r'\b(0?[1-9]|1[0-2])[-/](0?[1-9]|[12]\d|3[01])[-/](?:19|20)?\d{2}\b', "Date Format (MM/DD/YY or YYYY)", "Personal Data (PII)", "Date format (MM/DD/YYYY)."),
        (r'\b(192\.168(?:\.\d{1,3}){2}|10(?:\.\d{1,3}){3}|172\.(?:1[6-9]|2\d|3[01])(?:\.\d{1,3}){2}|127(?:\.\d{1,3}){3})\b', "Internal/Local IP", "Network", "Internal/localhost IP address."),
        (r'\b[0-9a-fA-F]{8}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{4}-[0-9a-fA-F]{12}\b', "UUID Format", "Identifier", "UUID format."),
        (r'\b(?:Project|Codename|Initiative)[ -_]([A-Z][a-zA-Z0-9]+(?:[ -_][A-Z][a-zA-Z0-9]+)*)\b', "Project Codename", "Intellectual Property", "Potential project codename."),
    ]
    regex_findings_temp = []
    if original_text:
        for item in regex_patterns:
            pattern, type, category, description = item[0], item[1], item[2], item[3]; flags = item[4] if len(item) > 4 else 0
            try:
                 for match in re.finditer(pattern, original_text, flags=flags):
                      value = match.group(0); start, end = match.start(), match.end()
                      if value and start < end: regex_findings_temp.append({"category": category, "type": type, "value": value, "description": description, "span": (start, end)})
            except re.error as e: print(f"Warning: Skipping invalid regex pattern: {pattern} - {e}")
            except IndexError as e: print(f"Warning: Index error during regex: {pattern} - {e}")
    print(f"DEBUG: Regex Findings Temp = {regex_findings_temp}")
    # 5. Add Findings in Priority Order
    print("\nDEBUG: Adding findings with overlap checks...")
    for finding in regex_findings_temp: add_finding(finding["category"], finding["type"], finding["value"], finding["description"], finding["span"][0], finding["span"][1])
    for finding in temp_keyword_findings: add_finding(finding["category"], finding["type"], finding["value"], finding["description"], finding["span"][0], finding["span"][1])
    for finding in potential_value_findings: add_finding(finding["category"], finding["type"], finding["value"], finding["description"], finding["span"][0], finding["span"][1])
    for finding in spacy_findings_temp: add_finding(finding["category"], finding["type"], finding["value"], finding["description"], finding["span"][0], finding["span"][1])
    print("DEBUG: Finished adding findings.\n")
    # --- End Finding Addition ---
    sorted_findings = sorted(findings, key=lambda f: (f['span'][0], -f['span'][1])); print(f"DEBUG: Final Findings List (Sorted) = {sorted_findings}")
    return {"found_issues": len(sorted_findings) > 0, "findings": sorted_findings}

# --- Helper function to read and parse audit log ---
def read_parse_audit_log():
    log_entries = []
    error = None
    try:
        if os.path.exists(LOG_FILE_PATH):
            with open(LOG_FILE_PATH, 'r', encoding='utf-8') as f:
                lines = f.readlines(); lines.reverse() # Show newest first
                for line_num, line in enumerate(lines, 1):
                    parts = line.strip().split(' | ')
                    if len(parts) == 7:
                        entry = {}
                        try:
                            entry['timestamp'] = parts[0].strip()
                            entry['user'] = parts[1].split(':', 1)[-1].strip() if ':' in parts[1] else parts[1].strip()
                            entry['model'] = parts[2].split(':', 1)[-1].strip() if ':' in parts[2] else parts[2].strip()
                            entry['findings'] = parts[3].split(':', 1)[-1].strip() if ':' in parts[3] else parts[3].strip()
                            entry['original_excerpt'] = parts[4].split(':', 1)[-1].strip() if ':' in parts[4] else parts[4].strip()
                            entry['redacted_excerpt'] = parts[5].split(':', 1)[-1].strip() if ':' in parts[5] else parts[5].strip()
                            entry['error_msg'] = parts[6].split(':', 1)[-1].strip() if ':' in parts[6] else parts[6].strip()
                            log_entries.append(entry)
                        except IndexError as e: print(f"Warning: Could not parse log line {line_num} due to missing parts: '{line.strip()}' - Error: {e}")
                        except Exception as parse_e: print(f"Warning: Could not parse log line {line_num}: '{line.strip()}' - Error: {parse_e}")
                    else: print(f"Warning: Skipping malformed log line {line_num} (incorrect parts: {len(parts)}): '{line.strip()}'")
        else: print("Audit log file not found.")
    except Exception as e: print(f"Error reading audit log file: {e}"); error = f"Could not read audit log file: {str(e)}"
    return log_entries, error

# --- Flask Routes ---
@app.route('/', methods=['GET', 'POST'])
def home():
    results = None; original_text_display = ""; uploaded_filename = None; error = None; llm_response = None; text_sent_to_llm = None; source_description = "No input processed yet."; placeholders_used = None; user_identifier = "anonymous"
    audit_log_entries = []; audit_log_error = None
    if request.method == 'POST':
        text_input_from_area = request.form.get('text_to_analyze', ''); file = request.files.get('file_to_analyze'); content_to_analyze = None
        if file and file.filename:
            if allowed_file(file.filename):
                uploaded_filename = secure_filename(file.filename); print(f"Processing uploaded file: {uploaded_filename}"); source_description = f"Uploaded file '{uploaded_filename}'"; file_ext = uploaded_filename.rsplit('.', 1)[1].lower()
                try:
                    if file_ext == 'txt': content_to_analyze = extract_text_from_txt(file.stream)
                    elif file_ext == 'pdf': content_to_analyze = extract_text_from_pdf(file.stream)
                    elif file_ext == 'docx': content_to_analyze = extract_text_from_docx(file.stream)
                    if content_to_analyze is None or content_to_analyze == "[ERROR: PDF is encrypted]":
                        error_reason = "Check file content/corruption/libraries."
                        if content_to_analyze == "[ERROR: PDF is encrypted]":
                            error_reason = "PDF is password-protected."
                        error = f"Failed to extract text from '{uploaded_filename}'. {error_reason}"
                        original_text_display = text_input_from_area
                        content_to_analyze = None
                    else:
                        original_text_display = ""
                except Exception as e: print(f"Error processing file {uploaded_filename}: {e}"); error = f"An error occurred processing file '{uploaded_filename}'."; original_text_display = text_input_from_area
            else: error = f"Invalid file type '{file.filename}'. Allowed: {', '.join(ALLOWED_EXTENSIONS)}."; original_text_display = text_input_from_area
        elif text_input_from_area: print("Processing text from input area."); content_to_analyze = text_input_from_area; original_text_display = text_input_from_area; source_description = "Text Input"
        else: error = "Please enter text or upload an allowed document (.txt, .pdf, .docx)."

        log_model = "N/A"; log_orig_excerpt = "N/A"; log_red_excerpt = "N/A"; log_num_findings = 0; log_error = error if error else ""
        if content_to_analyze is not None and error is None:
            if not isinstance(content_to_analyze, str): content_to_analyze = str(content_to_analyze)
            log_orig_excerpt = (content_to_analyze[:100] + '...') if len(content_to_analyze) > 100 else content_to_analyze
            print(f"Performing confidentiality check on content from {source_description}..."); results = perform_confidentiality_check(content_to_analyze); log_num_findings = len(results.get('findings', [])); print(f"Check complete. Found issues: {results['found_issues']}")
            if results['found_issues']:
                print("Redacting text before sending to LLM..."); used_placeholders_set = set(); findings_to_process = results.get('findings', [])
                if findings_to_process:
                    for finding in findings_to_process: used_placeholders_set.add(get_placeholder(finding))
                placeholders_used = sorted(list(used_placeholders_set)); text_sent_to_llm = redact_text(content_to_analyze, results['findings']); log_red_excerpt = (text_sent_to_llm[:100] + '...') if len(text_sent_to_llm) > 100 else text_sent_to_llm
            else: print("No issues found needing redaction, sending original text."); text_sent_to_llm = content_to_analyze; log_red_excerpt = log_orig_excerpt; placeholders_used = None
            try: log_model = API_URL.split('/models/')[-1] if '/models/' in API_URL else API_URL
            except Exception: log_model = "ErrorParsingURL"
            llm_response = query_llm(text_sent_to_llm)
            if llm_response and llm_response.startswith(("[Error", "[LLM Info")): log_error += f" LLM Query: {llm_response}"
            print("LLM query finished.")
        elif error: results = None; llm_response = None; text_sent_to_llm = None; placeholders_used = None; log_orig_excerpt = "Error processing input"; log_red_excerpt = "Error processing input"
        try: safe_orig_excerpt = log_orig_excerpt.replace('|', '/'); safe_red_excerpt = log_red_excerpt.replace('|', '/'); safe_error_msg = log_error.replace('|', '/'); log_message = (f"User:{user_identifier} | Model:{log_model} | Findings:{log_num_findings} | OriginalExcerpt:{safe_orig_excerpt} | RedactedExcerpt:{safe_red_excerpt} | Error:{safe_error_msg}"); audit_logger.info(log_message)
        except Exception as log_e: print(f"!!! FAILED TO WRITE AUDIT LOG: {log_e} !!!")

    # Read Audit Log on GET and POST requests
    audit_log_entries, audit_log_error = read_parse_audit_log()

    return render_template('index.html',
                           results=results, original_text=original_text_display,
                           uploaded_filename=uploaded_filename, error=error,
                           redacted_text=text_sent_to_llm, llm_response=llm_response,
                           placeholders_used=placeholders_used,
                           audit_log_entries=audit_log_entries, # Pass audit data
                           audit_log_error=audit_log_error)     # Pass audit error


# --- Main Execution ---
if __name__ == '__main__':
    print("-" * 50); print("Starting GenAI Shield Application...")
    print("Ensure required libraries are installed:"); print("  pip install Flask requests PyPDF2 python-docx python-dotenv spacy")
    print(f"Allowed file extensions: {ALLOWED_EXTENSIONS}")
    if not PyPDF2: print("  (PDF processing disabled)")
    if not docx: print("  (DOCX processing disabled)")
    if not nlp: print("  (spaCy NER processing disabled - install spacy and download model 'en_core_web_sm')")
    else: print("  (spaCy NER processing enabled)")
    print("Audit logs will be written to 'audit.log'"); print("-" * 50)
    if HF_API_KEY: print(f"Using Hugging Face API Key: YES"); print(f"Using LLM Endpoint: {API_URL}"); print("Ensure you have accepted terms for the selected model on huggingface.co (if applicable)")
    else: print("WARNING: Hugging Face API Key (HF_API_KEY) not found."); print("         LLM functionality will be disabled.")
    print("-" * 50); app.run(debug=True, host='0.0.0.0', port=5000)

# --- END OF FILE app.py ---